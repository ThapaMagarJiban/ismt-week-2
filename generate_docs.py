import os
from docx import Document
from docx.shared import Pt

plantuml_code = """
@startuml
abstract class BaseEntity {
    + int Id
    + DateTime CreatedAt
}

class Event {
    + string Name
    + DateTime Date
    + string Description
}

class Participant {
    + string Name
    + string Email
    + string Phone
}

class Venue {
    + string Name
    + string Address
    + int Capacity
}

class Activity {
    + string Name
    + string Type
}

class Registration {
    + DateTime RegistrationDate
    + string Status
}

class EventVenue {
}

class EventActivity {
}

BaseEntity <|-- Event
BaseEntity <|-- Participant
BaseEntity <|-- Venue
BaseEntity <|-- Activity

Event "1" -- "many" Registration
Participant "1" -- "many" Registration
(Event, Participant) .. Registration

Event "1" -- "many" EventVenue
Venue "1" -- "many" EventVenue

Event "1" -- "many" EventActivity
Activity "1" -- "many" EventActivity

interface IRepository<T> {
    + Task<IEnumerable<T>> GetAllAsync()
    + Task<T> GetByIdAsync(int id)
    + Task AddAsync(T entity)
    + Task UpdateAsync(T entity)
    + Task DeleteAsync(int id)
}

class Repository<T> {
}
IRepository <|.. Repository

interface IEventService {
    + Task CreateEventAsync(Event newEvent)
    + Task RegisterParticipantAsync(int eventId, int participantId)
}

class EventService {
}
IEventService <|.. EventService
EventService --> IRepository : uses
@enduml
"""

doc = Document()
doc.add_heading('System Documentation & Testing', 0)

doc.add_heading('1. UML Class Diagram', level=1)
doc.add_paragraph('Below is the PlantUML source for the Class Diagram illustrating the architecture of the Community Event Management System. It demonstrates inheritance (BaseEntity), interfaces (IRepository, IEventService), and associations (many-to-many relationships).')

p = doc.add_paragraph(plantuml_code)
p.style.font.name = 'Courier New'
p.style.font.size = Pt(9)

doc.add_heading('2. Test Documentation', level=1)
doc.add_paragraph('The application is tested using xUnit and Moq. Below are the test scenarios, methods, and outcomes.')

table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Test ID'
hdr_cells[1].text = 'Test Case Description'
hdr_cells[2].text = 'Test Method'
hdr_cells[3].text = 'Outcome'

test_cases = [
    ('TC-01', 'Verify CreateEventAsync calls Repository AddAsync once', 'Unit Test (Mocking)', 'Pass'),
    ('TC-02', 'Verify RegisterParticipantAsync successfully adds registration', 'Unit Test (InMemory DB)', 'Pass'),
    ('TC-03', 'Verify RegisterParticipantAsync throws exception if already registered', 'Unit Test (InMemory DB)', 'Pass')
]

for tc in test_cases:
    row_cells = table.add_row().cells
    row_cells[0].text = tc[0]
    row_cells[1].text = tc[1]
    row_cells[2].text = tc[2]
    row_cells[3].text = tc[3]

doc.add_paragraph('\nThese tests validate the service layer, particularly the `EventService`, handling core business logic for event creation and participant registration.')

doc.save('Documentation.docx')
print("Documentation generated.")
